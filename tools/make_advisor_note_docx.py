"""Builds docs/VRDS_Experiment2_Research_Questions_and_Hypotheses.docx - a short, simple note from Waseem to his
advisor: the research questions, the hypotheses, what the main scores mean, and an empty table for suggestions.
Content comes from docs/ANALYSIS_AND_FIGURES_PLAN.md and docs/PREREGISTRATION_FREEZE.md (27 September 2026)."""
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Cm

OUT = sys.argv[1]
BLUE = RGBColor(0x1F, 0x38, 0x64)      # headings
TEXT = RGBColor(0x26, 0x26, 0x26)      # body
GREY = RGBColor(0x59, 0x59, 0x59)      # subtitle, notes
HEADER_FILL = "DCE6F1"                 # table header row
GROUP_FILL = "F2F2F2"                  # group label rows
LINE = "BFBFBF"                        # table lines

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.9))
sec.top_margin, sec.bottom_margin = Inches(0.8), Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.font.size = Pt(10.5)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(4)


def para(text="", size=10.5, bold=False, color=TEXT, italic=False, after=4, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.italic = Pt(size), bold, italic
        r.font.color.rgb = color
    return p


def heading(text):
    return para(text, size=14, bold=True, color=BLUE, before=14, after=6)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table):
    tbl_pr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), LINE)
        b.append(e)
    tbl_pr.append(b)


def cell_text(cell, text, bold=False, size=10, color=TEXT, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(1)
    r = p.add_run(text)
    r.font.size, r.bold, r.italic = Pt(size), bold, italic
    r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    tr_pr.append(h)


def no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    c = OxmlElement("w:cantSplit")
    c.set(qn("w:val"), "true")
    tr_pr.append(c)


def table(headers, rows, widths):
    """rows: lists of cell strings, or ("GROUP", label) for a shaded group row spanning the table."""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        cell_text(c, h, bold=True, color=BLUE)
        shade(c, HEADER_FILL)
    repeat_header(t.rows[0])
    for row in rows:
        new_row = t.add_row()
        no_split(new_row)
        cells = new_row.cells
        if isinstance(row, tuple) and row[0] == "GROUP":
            merged = cells[0].merge(cells[-1])
            cell_text(merged, row[1], bold=True, color=BLUE, size=10)
            shade(merged, GROUP_FILL)
            for par in merged.paragraphs:
                par.paragraph_format.keep_with_next = True  # a group label never ends a page alone
            continue
        for i, v in enumerate(row):
            cell_text(cells[i], v, bold=(i == 0))
    for i, w in enumerate(widths):
        t.columns[i].width = Inches(w)
    for r in t.rows:
        if len({id(c._tc) for c in r.cells}) == 1:
            continue  # a merged group row keeps its full width
        for i, w in enumerate(widths):
            r.cells[i].width = Inches(w)
    return t


# ---------------------------------------------------------------- title
para("VRDS Experiment 2", size=22, bold=True, color=BLUE, after=0)
para("Research questions, hypotheses and measures for the analysis phase", size=12.5, color=GREY, after=2)
para("Waseem Samkari  ·  September 2026", size=10, color=GREY, after=10)
para("This note summarizes what I plan to learn from the Experiment 2 data: the research questions, the hypotheses "
     "I will test, and what each main score means. Participants build a value profile in Blocks 1–4, then make four "
     "decisions and one wish in emergency scenarios where the cost lands on different people, and finish with a "
     "prediction test and a feedback questionnaire. The last page is for your suggestions.", after=6)

# ---------------------------------------------------------------- 1. research questions
heading("1. Research questions")
W3 = [0.45, 4.35, 1.9]
table(["#", "Research question", "Answered with"], [
    ("GROUP", "Position: who carries the cost"),
    ["Q1", "Do people keep the same moral priorities when the cost lands on someone else?", "Distance of each choice from the participant's values, in five positions"],
    ["Q2", "Deciding for colleagues vs. having the same decision done to you: do the choices differ?", "Scenarios 4 and 5 (a matched pair)"],
    ["Q3", "Which value rises when a decision lands on the participant?", "The wish minus the decision, per value"],
    ["Q4", "Under an employer's values, do people adopt them, compromise, or hold their own?", "Company stance"],
    ["Q5", "Do people cross their own red lines more when others carry the cost?", "Red-line crossings by position"],
    ("GROUP", "Reflection (CVR) and clarification (APA)"),
    ["Q6", "How often do people choose against their own values, and what do they do after the reflection?", "CVR outcomes"],
    ["Q7", "When the reflection changes a choice, does the new choice fit their values better?", "Fit before vs. after"],
    ["Q8", "Does reflection cost performance?", "Performance before vs. after"],
    ["Q9", "Does the value named in the clarification guide the next decisions?", "APA value vs. later choices"],
    ["Q10", "Does reading the MCF go with different choices?", "MCF reading vs. the choice"],
    ("GROUP", "Values and performance"),
    ["Q11", "How do people trade their own values against performance?", "VCI vs. performance"],
    ["Q12", "Does the Blocks 1–4 profile predict Block 5 choices beyond card order and performance?", "A choice model"],
    ("GROUP", "Prediction (MPF)"),
    ["Q13", "Can the MPF predict choices better than chance, and are its percentages accurate?", "Hit rate and calibration"],
    ["Q14", "Does learning during Block 5 improve the prediction?", "Learning vs. fixed predictor"],
    ["Q15", "Behind the veil, does the chosen rule match the values people act on, and do they react to being predicted?", "Scenario 6"],
    ("GROUP", "Blocks 1–4"),
    ["Q16", "Do Blocks 1–3 reproduce known effects (push vs. pull, the shelter, vulnerable workers)?", "Blocks 1–3 thresholds"],
    ["Q17", "Does context sensitivity with money carry over to decisions about lives?", "Block 1 vs. position effect"],
    ["Q18", "Is reconsidering a personal trait, from Block 4 to Block 5?", "Block 4 change vs. Block 5 change"],
    ("GROUP", "Experience (feedback)"),
    ["Q19", "Are consistent and stable participants more satisfied, and are high performers less satisfied?", "Well-being vs. the main scores"],
    ["Q20", "Do participants' self-reports match their behavior?", "E.g., value congruence vs. VCI"],
    ["Q21", "Do participants know which value matters most to them?", "Open answer vs. their top value"],
], W3)

# ---------------------------------------------------------------- 2. hypotheses
heading("2. Hypotheses I will test")
para("All hypotheses are fixed before I open the data.", size=10, color=GREY, after=4)
table(["", "Hypothesis", "Test"], [
    ("GROUP", "Primary"),
    ["H1", "Departure from one's own values differs between deciding for colleagues (scenario 4) and having the same decision done to oneself (scenario 5).", "Wilcoxon signed-rank"],
    ["H2", "When the reflection changes a choice, the new choice fits the participant's values better.", "Paired Wilcoxon, one-sided"],
    ["H3", "Reflection does not cost performance (within ±5 points).", "Equivalence test (TOST)"],
    ["H4", "The MPF predicts the first choice better than chance (16.7%).", "Bootstrap by participant"],
    ("GROUP", "Secondary"),
    ["H5", "Consistency (VCI) is above blind choice (50) and above random choosing (57).", "One-sample Wilcoxon"],
    ["H6", "The Blocks 1–4 values predict Block 5 choices beyond card position and performance.", "Conditional logit"],
    ["H7", "The learning MPF predicts better than a version that never learns.", "Paired comparison"],
    ["H8", "Well-being and satisfaction rise with VCI and Stability, and fall with performance.", "Multiple regression"],
    ["H9", "Self-reported value congruence rises with VCI.", "Spearman correlation"],
    ["H10", "People whose money decisions change more across places also change more across positions in Block 5.", "Spearman, controlling for VCI"],
    ["H11", "People who changed their Block 4 decision also change their choice more often after the Block 5 reflection.", "Mixed logistic regression"],
], W3)
para("Primary hypotheses are corrected together (Holm); secondary ones with a false-discovery-rate correction.",
     size=9.5, color=GREY, before=4, after=4)

# ---------------------------------------------------------------- 3. scores
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
heading("3. The main scores and what they mean")
table(["Score", "What it means", "How to read it"], [
    ["Value profile (seven sensitivities)", "How much each of seven values matters to the participant, measured in Blocks 1–4: protecting the vulnerable, reducing harm, how many are helped, how much is gained, directness, context, and how much others' perspectives move them.", "0–100 for each value"],
    ["VCI (Value Consistency Index)", "How well the four decisions fit the participant's own values.", "50 = blind choice · 100 = always the best fit"],
    ["Stability", "Whether the order of the participant's values holds when a decision goes against them.", "100 = the order held"],
    ["Performance", "How much of the best available outcome the decisions captured.", "0 = the weakest option · 100 = the strongest"],
    ["Position Effect", "How much choices shift with who carries the cost: oneself, one's family, strangers, colleagues under an employer, or oneself on the receiving end.", "Higher = position mattered more"],
    ["CVR (value-reflection step)", "When a choice goes against the participant's values, it shows the consequence from two views and the affected person's voice, then asks them to keep or reconsider.", "Kept or reconsidered"],
    ["APA (value-clarification step)", "Lets the participant say which value they want weighted; the profile follows their answer.", "The value named, with confidence 1–5"],
    ["MCF (Moral Commitment Function)", "For each option, says in plain words what it gives and asks of the participant's own values, and which option serves each value most.", "Read or not, and what was read"],
    ["MPF (Moral Prediction Function)", "Predicts the participant's choice from their values, as a percentage for each option, and is tested openly in scenario 6.", "Chance = 16.7% (six options) or 25% (four)"],
], [1.75, 3.55, 1.4])

# ---------------------------------------------------------------- 4. suggestions (own page)
p = doc.add_paragraph()
p.add_run().add_break(WD_BREAK.PAGE)
heading("4. Your suggestions")
para("I would welcome your ideas on any analysis, measure or data I should add or collect.", after=6)
t = table(["#", "What should I add or collect?", "Why it would help", "Priority"],
          [[str(i), "", "", ""] for i in range(1, 9)], [0.4, 2.9, 2.6, 0.8])
for r in t.rows[1:]:
    r.height = Cm(1.5)
para("Thank you for your time.", before=12, after=0)
para("Waseem", bold=True, after=0)

doc.core_properties.title = "VRDS Experiment 2 - Research questions, hypotheses and measures"
doc.core_properties.author = "Waseem Samkari"
doc.save(OUT)
print("saved", OUT)
