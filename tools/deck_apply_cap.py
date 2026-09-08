"""
deck_apply_cap.py — brings the APA slides in line with the per-value cap.

The rule now carries a cap: no policy value moves more than 30 x confidence in one clarification.
That changes what several slides should say, so this rewrites them to describe the mechanism as it
stands — not as a sequence of changes.

  · the rules table gains the cap row
  · "Problem 1" becomes "Why there is a cap"
  · "Problems 2 and 3" becomes what the cap protects, and what remains weak
  · the persona table's most extreme row now stops short of the ceiling
  · the rating is restated
  · the "two proposed fixes" slide goes: the remaining open item is one line on the weakness slide

Refuses to run twice. Both decks, one definition.

Usage:  python tools/deck_apply_cap.py [--ar]
"""
import sys, os, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
import deck_style as D
from deck_style import head, box, table, blank, MARGIN_L, CONTENT_W, PANEL, BLUE, ORANGE, GREEN, TEAL

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "Generated Outputs")
IN = lambda v: int(v * 914400)
L, W = MARGIN_L, CONTENT_W

MARK_EN, MARK_AR = "Why there is a cap", "لماذا يوجد سقف"
CAP_ROW_EN = ["The cap", "no policy value moves more than 30 × that scale, in either direction"]
CAP_ROW_AR = ["السقف", "لا تتحرّك أي قيمة أكثر من 30 × ذلك المعامل، في أي اتجاه"]
DROP_EN, DROP_AR = "Two proposed fixes", "إصلاحان مقترحان"

FIXES_EN = [
    ("gained 79→100 (+21)", "gained 79→97 (+18)"),
    ("Sound mechanism, with one interaction left to decide.", "Two weak points, both stated."),
    ("Works in isolation, defeated on the endorse path",
     "Works — the largest possible move rises with the rating and nothing else"),
    ('True for "just this situation". FALSE for the endorse path.',
     "The published constant is the applied one, on every answer path."),
    ("One persona reaches the ceiling. The weakest part of the mechanism.",
     "11.9% of values reach 100. The weakest part of the mechanism."),
    ("High — the fix is small and local to one function",
     "Six are set by judgement; the stakeholder ±25 has no measurement behind it"),
    ("Fixability", "Constants"),
    ("8 / 10", "9 / 10"),
]
FIXES_AR = [
    ("المكسب 79←100 (+21)", "المكسب 79←97 (+18)"),
    ("آلية سليمة، مع تفاعل واحد بانتظار قرارك.", "نقطتا ضعف، وكلتاهما مذكورة."),
    ("تعمل منفردة، ويُبطلها الاحتساب المزدوج في مسار التأييد",
     "تعمل — وأكبر حركة ممكنة ترتفع مع الدرجة ولا شيء غيرها"),
    ("صحيحة في مسار «لهذا الموقف». وغير صحيحة في مسار التأييد.",
     "الثابت المنشور هو المطبَّق فعلاً، في كل المسارات."),
    ("واحد يبلغ السقف. وهو أضعف جزء في الآلية.", "‏11.9% من القيم تبلغ 100. وهو أضعف جزء في الآلية."),
    ("عالية — الإصلاح صغير ومحصور في دالة واحدة",
     "ستة ثوابت بالاجتهاد؛ و±25 لأصحاب المصلحة بلا قياس خلفه"),
    ("قابلية الإصلاح", "الثوابت"),
    ("8 / 10", "9 / 10"),
]


def add_cap_row(prs, ar):
    """Insert the cap into the rules table, directly after the confidence row."""
    key = "confidence 1" if not ar else "الثقة من 1"
    cells_text = CAP_ROW_AR if ar else CAP_ROW_EN
    if ar:
        cells_text = list(reversed(cells_text))
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_table:
                continue
            rows = list(sh.table.rows)
            for i, row in enumerate(rows):
                joined = " ".join(c.text for c in row.cells)
                if key not in joined:
                    continue
                new_tr = copy.deepcopy(row._tr)
                row._tr.addnext(new_tr)
                fresh = list(sh.table.rows)[i + 1]
                for cell, txt in zip(fresh.cells, cells_text):
                    tf = cell.text_frame
                    for extra in list(tf.paragraphs)[1:]:
                        extra._p.getparent().remove(extra._p)
                    p0 = tf.paragraphs[0]
                    if p0.runs:
                        p0.runs[0].text = txt
                        for r in list(p0.runs)[1:]:
                            r._r.getparent().remove(r._r)
                        p0.runs[0].font.bold = True
                return True
    return False


def rebuild_problem1(prs, ar):
    """'Problem 1' becomes the reason the cap exists."""
    title_key = "Problem 1" if not ar else "المشكلة 1"
    for idx, s in enumerate(prs.slides):
        heads = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if not heads or title_key not in heads[0]:
            continue
        page = None
        for sh in s.shapes:
            if (sh.has_text_frame and sh.left is not None
                    and abs(sh.top - D.PAGENO[1]) < 20000 and sh.text_frame.text.strip().isdigit()):
                page = sh.text_frame.text.strip()
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        if ar:
            head(s, MARK_AR, "سؤالان قد يسمّيان القيمة نفسها. ولولا السقف لتجمّعا.", page or 0)
            box(s, L, IN(1.80), W, IN(0.78),
                "لا شيء يمنع زيادة السؤال الأول وزيادة السؤال الثاني من الوقوع على القيمة نفسها.",
                fill=ORANGE, size=16, bold=True)
            table(s, L, IN(2.80), W, [
                ["نمط الإجابة", "الأول", "الثاني", "لولا السقف"],
                ["يؤيّد الخيار ويسمّي القيمة التي خدمها", "‏+15 × w", "‏+30 × w", "‏+45 × w"],
            ], [0.52, 0.16, 0.16, 0.16], row_h=[IN(0.45), IN(0.62)], size=14)
            box(s, L, IN(4.10), W, IN(1.45),
                "لماذا هذه هي الحالة الشائعة لا الاستثنائية\n"
                "من يدافع عن اختياره يسمّي بطبيعة الحال القيمة التي حماها ذلك الاختيار.\n"
                "فلولا السقف لكان الثابت المنشور خاطئاً لدى أكثر المشاركين اتساقاً، لا لدى أكثرهم ارتباكاً.",
                fill=PANEL, size=14)
            box(s, L, IN(5.75), W, IN(0.85),
                "السقف يجعل «+30 × الثقة» أكبر حركة ممكنة فعلاً، لكل مشارك وفي كل مسار إجابة.\n"
                "وتتحقّق منه البوابة A-APA-8 عبر كل تركيبات الإجابات ومستويات الثقة الخمسة.",
                fill=GREEN, size=14, bold=True)
        else:
            head(s, MARK_EN, "Two questions can name the same value. Without a cap they would add.", page or 0)
            box(s, L, IN(1.80), W, IN(0.78),
                "Nothing stops Q1's bump and Q2's bump landing on the SAME value.",
                fill=ORANGE, size=16, bold=True)
            table(s, L, IN(2.80), W, [
                ["Answer pattern", "Q1", "Q2", "Without the cap"],
                ["Endorses the option AND names the value it served", "+15 × w", "+30 × w", "+45 × w"],
            ], [0.52, 0.16, 0.16, 0.16], row_h=[IN(0.45), IN(0.62)], size=14)
            box(s, L, IN(4.10), W, IN(1.45),
                "WHY THIS IS THE COMMON CASE, NOT AN EDGE CASE\n"
                "Someone who defends a choice naturally goes on to name the value that choice protected.\n"
                "So without the cap, the published constant would be wrong for the most CONSISTENT participants,\n"
                "not the most confused ones.",
                fill=PANEL, size=14)
            box(s, L, IN(5.75), W, IN(0.85),
                "The cap makes \"+30 × confidence\" the true maximum for every participant, on every answer path.\n"
                "Gate A-APA-8 asserts it across every answer combination and all five confidence levels.",
                fill=GREEN, size=14, bold=True)
        return True
    return False


def rebuild_problem23(prs, ar):
    """'Problems 2 and 3' becomes what the cap protects, and what stays weak."""
    title_key = "Problems 2 and 3" if not ar else "المشكلتان 2 و3"
    for s in prs.slides:
        heads = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if not heads or title_key not in heads[0]:
            continue
        page = None
        for sh in s.shapes:
            if (sh.has_text_frame and sh.left is not None
                    and abs(sh.top - D.PAGENO[1]) < 20000 and sh.text_frame.text.strip().isdigit()):
                page = sh.text_frame.text.strip()
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        if ar:
            head(s, "ما يحميه السقف، وما يبقى ضعيفاً", "الآلية سليمة. ونقطتا الضعف مذكورتان لا مخفيّتان.", page or 0)
            box(s, L, IN(1.80), W, IN(0.55), "ما يحميه — درجة الثقة", fill=GREEN, size=16, bold=True)
            box(s, L, IN(2.45), W, IN(1.15),
                "لولا السقف، لتحرّك المشارك المحتسَب مرتين — من الثقة 2 فصاعداً — أبعد من الواثق تماماً:\n"
                "‏45 × 0.7 = 31.5 مقابل 30 × 1.0. ومع السقف تصبح أكبر حركة 30 × w، ترتفع مع الدرجة ولا شيء غيرها.",
                fill=PANEL, size=14)
            box(s, L, IN(3.85), W, IN(0.55), "ما يبقى ضعيفاً", fill=ORANGE, size=16, bold=True)
            box(s, L, IN(4.50), W, IN(1.65),
                "الالتصاق بالسقف: 11.9% من القيم تبلغ 100 بعد التوضيح. والقيمة الملتصقة لا تُظهر تأييداً إضافياً،\n"
                "ويفقد الترتيب دقّته عند القمة حيث تُقرَّر تسميات التوافق. وهذا ملازم للزيادات الثابتة على مقياس محدود.\n\n"
                "و±25 لأصحاب المصلحة: كبيرة، وغير مضروبة في الثقة، وليس خلف مقدارها أي قياس.",
                fill=PANEL, size=14)
        else:
            head(s, "What the cap protects, and what stays weak",
                 "The mechanism is sound. The two weak points are stated, not hidden.", page or 0)
            box(s, L, IN(1.80), W, IN(0.55), "WHAT IT PROTECTS — the confidence rating", fill=GREEN, size=16, bold=True)
            box(s, L, IN(2.45), W, IN(1.15),
                "Uncapped, from confidence 2 upward a double-counting participant moves FURTHER than one who is\n"
                "completely sure and does not: 45 × 0.7 = 31.5 against 30 × 1.0.\n"
                "With the cap the largest possible move is 30 × w — it rises with the rating and nothing else.",
                fill=PANEL, size=14)
            box(s, L, IN(3.85), W, IN(0.55), "WHAT STAYS WEAK", fill=ORANGE, size=16, bold=True)
            box(s, L, IN(4.50), W, IN(1.65),
                "SATURATION — 11.9% of values reach 100 after a clarification. A pinned value cannot show further\n"
                "endorsement, and the ranking loses resolution at the top, where alignment labels are decided.\n"
                "Inherent to flat deltas on a bounded scale: the cap reduces it, nothing removes it.\n\n"
                "THE STAKEHOLDER ±25 — large, unscaled, and the one constant with no measurement behind it.",
                fill=PANEL, size=14)
        return True
    return False


def drop_slide(prs, marker):
    for i, s in enumerate(prs.slides):
        heads = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if heads and marker in heads[0]:
            ids = prs.slides._sldIdLst
            entry = list(ids)[i]
            prs.part.drop_rel(entry.rId)
            ids.remove(entry)
            return True
    return False


def retext(prs, pairs):
    n = 0
    def walk(tf):
        nonlocal n
        for para in tf.paragraphs:
            for r in para.runs:
                for a, b in pairs:
                    if a in r.text:
                        r.text = r.text.replace(a, b); n += 1
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                walk(sh.text_frame)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        walk(c.text_frame)
    return n


def renumber(prs, ar):
    x = D.PAGENO_RTL[0] if ar else D.PAGENO[0]
    n = 0
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if (sh.has_text_frame and sh.left is not None
                    and abs(sh.left - x) < 20000 and abs(sh.top - D.PAGENO[1]) < 20000):
                if sh.text_frame.paragraphs and sh.text_frame.paragraphs[0].runs:
                    sh.text_frame.paragraphs[0].runs[0].text = str(i); n += 1
                break
    return n


def main(ar):
    name = "Block5_Design_Update_ARABIC" if ar else "Block5_Design_Update_for_Advisor"
    path = os.path.join(OUT, name + ".pptx")
    D.set_locale(font="Arial" if ar else "Calibri", rtl=ar)
    prs = Presentation(path)
    marker = MARK_AR if ar else MARK_EN
    if any(marker in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame):
        print("REFUSING: %s already carries the cap slides." % os.path.basename(path))
        return 1

    before = len(prs.slides)
    caprow = add_cap_row(prs, ar)
    p1 = rebuild_problem1(prs, ar)
    p23 = rebuild_problem23(prs, ar)
    swaps = retext(prs, FIXES_AR if ar else FIXES_EN)
    dropped = drop_slide(prs, DROP_AR if ar else DROP_EN)
    pages = renumber(prs, ar)
    prs.save(path)
    print("OK  %s" % os.path.basename(path))
    print("    %d -> %d slides  ·  cap row: %s  ·  'why a cap': %s  ·  'what stays weak': %s  ·  swaps: %d  ·  fixes slide dropped: %s  ·  pages: %d"
          % (before, len(prs.slides), caprow, p1, p23, swaps, dropped, pages))
    return 0


if __name__ == "__main__":
    sys.exit(main("--ar" in sys.argv))
