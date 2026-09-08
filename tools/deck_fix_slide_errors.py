"""
deck_fix_slide_errors.py — corrects four errors on the APA slides.

WHAT WAS WRONG
--------------
1. AN ARITHMETIC ERROR, and the serious one. Slide 27 claimed "+45 x 0.6 = +27, which is larger
   than a fully confident participant's +30 x 1.0". It is not: 27 is smaller than 30. The point
   being made is real but it starts at confidence 2, not 1 — 45 x 0.7 = 31.5, which does beat 30.
   Persona B (confidence 1) still reached 100/100, but because they started at 79 and hit the
   ceiling, not because 27 beats 30.

2. "un-add togethered" — produced by a blind word swap of "stack" to "add together" on the earlier
   restructure pass. Appears twice.

3. "for one follow-up questions" — the same kind of damage, from "clarification".

4. The Stability churn ceiling is quoted as 56. It moved to 57 on 7 Sept 2026.

Refuses to run twice. Both decks, one definition.

Usage:  python tools/deck_fix_slide_errors.py [--ar]
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
import deck_style as D

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "Generated Outputs")

FIX_EN = [
    # 1 — the arithmetic
    ("Persona B answered confidence 1 — NOT SURE — and finished that value at 100 / 100.",
     "Persona B answered confidence 1 — NOT SURE — and still finished that value at 100 / 100,"
     " because they began at 79 and ran into the ceiling."),
    ("+45 × 0.6  =  +27, which is larger than a fully confident participant's un-add togethered +30 × 1.0.",
     "Above confidence 1 it gets worse: +45 × 0.7  =  +31.5, which beats a completely sure participant's +30 × 1.0."),
    ("The least sure answer produced a bigger move than the most sure one.",
     "From confidence 2 upward, being LESS sure can move the model FURTHER than being completely sure."),
    # 2 and 3 — word-swap damage
    ("Both fixes leave the un-add togethered path alone", "Both fixes leave the un-doubled path alone"),
    ("Clamp each value's NET change for one follow-up questions to ±30 × confidence.",
     "Clamp each value's NET change, for one set of follow-up questions, to ±30 × confidence."),
    # 4 — the ceiling
    ("Stability's churn ceiling (56) was calibrated on current APA behaviour.",
     "Stability's churn ceiling (57) is calibrated on current APA behaviour."),
]

FIX_AR = [
    ("المشارك B أجاب بثقة 1 — أي «لستُ متأكداً» — وانتهى بتلك القيمة عند 100 من 100.",
     "المشارك B أجاب بثقة 1 — أي «لستُ متأكداً» — وانتهى مع ذلك عند 100 من 100، لأنه بدأ من 79 فبلغ السقف."),
    ("‏+45 × 0.6  =  +27، وهي أكبر من +30 × 1.0 لمشارك واثق تماماً لم يحدث لديه تراكم.",
     "وفوق الثقة 1 يسوء الأمر: +45 × 0.7  =  +31.5، وهي تتجاوز +30 × 1.0 لمشارك واثق تماماً."),
    ("أي أن الإجابة الأقل تأكداً أنتجت حركة أكبر من الإجابة الأكثر تأكداً.",
     "من الثقة 2 فصاعداً، قد يحرّك الأقلُّ تأكداً النموذجَ أبعد من الواثق تماماً."),
    # narrower than the English pattern: the sentence around it differs, and "عوير" was also a typo
    ("سقف Stability (56) عُوير", "سقف Stability (57) مُعايَر"),
]


def apply(prs, pairs):
    hits = []
    def walk(tf, where):
        for para in tf.paragraphs:
            for r in para.runs:
                for a, b in pairs:
                    if a in r.text:
                        r.text = r.text.replace(a, b)
                        hits.append((where, a[:44]))
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame:
                walk(sh.text_frame, i)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        walk(c.text_frame, i)
    return hits


def main(ar):
    name = "Block5_Design_Update_ARABIC" if ar else "Block5_Design_Update_for_Advisor"
    path = os.path.join(OUT, name + ".pptx")
    D.set_locale(font="Arial" if ar else "Calibri", rtl=ar)
    prs = Presentation(path)
    pairs = FIX_AR if ar else FIX_EN
    hits = apply(prs, pairs)
    if not hits:
        print("REFUSING: %s — nothing matched. Already fixed?" % os.path.basename(path))
        return 1
    prs.save(path)
    print("OK  %s   (%d of %d corrections applied)" % (os.path.basename(path), len(hits), len(pairs)))
    for where, what in hits:
        print("      slide %-3d %s" % (where, what))
    return 0


if __name__ == "__main__":
    sys.exit(main("--ar" in sys.argv))
