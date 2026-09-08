"""
deck_position_five.py — the Position Effect slides, for the five-position deck.

Three slides still described a three-position ladder:

  · the headline formula said "highest of the three minus lowest of the three"
  · the two-units slide said the share exists so "the three positions" compare fairly
  · the separation table carried three position columns and four rows of figures

The deck now varies FIVE positions, one scenario each, and scenario 5's wish counts in this
measure even though it is excluded from VCI, CVR, APA and Stability. The separation table is
rebuilt from `npm run validate:position` on the shipping deck rather than kept from a smaller one,
which also means it now shows the honest result: the drifter and the random responder no longer
separate, because no position repeats and the drift check returns null.

Refuses to run twice. Both decks, one definition.

Usage:  python tools/deck_position_five.py [--ar]
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
import deck_style as D
from deck_style import head, box, table, MARGIN_L, CONTENT_W, PANEL, BLUE, ORANGE, GREEN

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "Generated Outputs")
IN = lambda v: int(v * 914400)
L, W = MARGIN_L, CONTENT_W

SENTINEL_EN = "across five profile types"
SENTINEL_AR = "على خمسة أنماط"

SWAP_EN = [
    ("Position Effect  =  highest of the three   minus   lowest of the three        Range 0–100.",
     "Position Effect  =  highest position   minus   lowest position        Five positions, one scenario each. Range 0–100."),
    ("The share removes that, so the three positions can be compared fairly.",
     "The share removes that, so all five positions can be compared fairly."),
    ("Points can be compared across\nall five scenarios.", "Points can be compared across\nall five scenarios."),
]
SWAP_AR = [
    ("أثر الموقع  =  أعلى المواقع الثلاثة  −  أدناها          المدى 0–100",
     "أثر الموقع  =  أعلى موقع  −  أدنى موقع        خمسة مواقع، سيناريو لكلٍّ منها. المدى 0–100"),
    ("وتزيل النسبة هذا الفرق، فتصبح المواقع الثلاثة قابلة للمقارنة بإنصاف.",
     "وتزيل النسبة هذا الفرق، فتصبح المواقع الخمسة قابلة للمقارنة بإنصاف."),
    ("والنسبة تزيل هذا الفرق، فتصبح المواقع الثلاثة قابلة للمقارنة بإنصاف.",
     "والنسبة تزيل هذا الفرق، فتصبح المواقع الخمسة قابلة للمقارنة بإنصاف."),
]


def swap(prs, pairs):
    n = 0
    def walk(tf):
        nonlocal n
        whole = tf.text
        for a, b in pairs:
            if a not in whole or a == b:
                continue
            if "\n" in a:
                continue
            for para in tf.paragraphs:
                for r in para.runs:
                    if a in r.text:
                        r.text = r.text.replace(a, b); n += 1
            whole = tf.text
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                walk(sh.text_frame)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        walk(c.text_frame)
    return n


def rebuild_separation(prs, ar):
    """The separation slide, re-measured on the shipping deck."""
    key = "does it separate people" if not ar else "هل يفرّق بين الناس"
    for s in prs.slides:
        heads = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if not heads or key not in heads[0]:
            continue
        page = None
        for sh in s.shapes:
            if (sh.has_text_frame and sh.left is not None
                    and abs(sh.top - D.PAGENO[1]) < 20000 and sh.text_frame.text.strip().isdigit()):
                page = sh.text_frame.text.strip()
        title = heads[0]
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        if ar:
            head(s, title, "محاكاة على مجموعة الخيارات الحقيقية، على خمسة أنماط من الملفات. وهذه البوابات تعمل مع كل بناء.", page or 0)
            table(s, L + IN(1.2), IN(1.90), W - IN(2.4), [
                ["نوع المشارك", "أثر الموقع", "فحص الانحراف"],
                ["الشخص نفسه في كل موقع", "0", "غير متاح"],
                ["يبدّل بحسب الموقع", "100", "غير متاح"],
                ["يختار بلا نمط", "100", "غير متاح"],
            ], [0.5, 0.25, 0.25], row_h=[IN(0.46)] + [IN(0.52)] * 3, size=14)
            box(s, L, IN(3.90), W, IN(0.80),
                "الصفّ الأول ينفصل تماماً — فارق 100 نقطة على كل نمط جُرِّب.\n"
                "أما الصفّان الثاني والثالث فلا ينفصلان عن بعضهما.",
                fill=ORANGE, size=15, bold=True)
            box(s, L, IN(4.90), W, IN(1.55),
                "فحص الانحراف وحده يفرّق بينهما، وهو يحتاج موقعاً يتكرّر في سيناريوهين. وفي هذا الترتيب يظهر كل موقع\n"
                "من المواقع الخمسة مرة واحدة، فيعود الفحص فارغاً. لذا فأثر الموقع وصفٌ لما فعله المشارك،\n"
                "لا دليلاً على أن الموقع هو السبب.\n"
                "وزوج مكان العمل (4 و5) مضبوط المحتوى — الشركة نفسها والقرار نفسه والخيارات الستة نفسها —\n"
                "فاستعمله لأي ادّعاء استدلالي عن الموقع.",
                fill=PANEL, size=13)
        else:
            head(s, title, "Simulated on the real option set, across five profile types. These gates run on every build.", page or 0)
            table(s, L + IN(1.2), IN(1.90), W - IN(2.4), [
                ["Participant type", "Position Effect", "Drift check"],
                ["Same person everywhere", "0", "unavailable"],
                ["Switches by position", "100", "unavailable"],
                ["Chooses without a pattern", "100", "unavailable"],
            ], [0.5, 0.25, 0.25], row_h=[IN(0.46)] + [IN(0.52)] * 3, size=14)
            box(s, L, IN(3.90), W, IN(0.80),
                "The first row separates cleanly — a 100-point gap on every profile tested.\n"
                "The second and third do NOT separate from each other.",
                fill=ORANGE, size=15, bold=True)
            box(s, L, IN(4.90), W, IN(1.55),
                "Only the drift check tells those two apart, and it needs one position to appear in two scenarios.\n"
                "Here each of the five appears once, so it returns null. Position Effect is therefore DESCRIPTIVE:\n"
                "what a participant did, not evidence that position caused it.\n"
                "The workplace pair (4 and 5) is content-controlled — same employer, same decision, same six options —\n"
                "so use that contrast for any inferential claim about position.",
                fill=PANEL, size=13)
        return True
    return False


def main(ar):
    name = "Block5_Design_Update_ARABIC" if ar else "Block5_Design_Update_for_Advisor"
    path = os.path.join(OUT, name + ".pptx")
    D.set_locale(font="Arial" if ar else "Calibri", rtl=ar)
    prs = Presentation(path)
    sentinel = SENTINEL_AR if ar else SENTINEL_EN
    if any(sentinel in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame):
        print("REFUSING: %s already updated." % os.path.basename(path))
        return 1
    n = swap(prs, SWAP_AR if ar else SWAP_EN)
    rebuilt = rebuild_separation(prs, ar)
    prs.save(path)
    print("OK  %s   ·   text swaps: %d   ·   separation slide rebuilt: %s"
          % (os.path.basename(path), n, rebuilt))
    return 0


if __name__ == "__main__":
    sys.exit(main("--ar" in sys.argv))
