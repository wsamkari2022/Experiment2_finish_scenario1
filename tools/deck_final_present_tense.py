"""
deck_final_present_tense.py — the last four places that still narrated a change.

Everything here describes what the system does, not what it used to do:

  slide 2   the summary line and its APA row
  slide 20  the Part 3 opener, which framed the rule as correct-but-incomplete
  slide 35  a stability-ceiling figure quoted from an older measurement
  slide 52  the feature table, whose "why" column explained earlier versions

Refuses to run twice. Both decks, one definition.

Usage:  python tools/deck_final_present_tense.py [--ar]
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx import Presentation
import deck_style as D
from deck_style import head, box, table, MARGIN_L, CONTENT_W, PANEL, BLUE, ORANGE, GREEN

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "Generated Outputs")
IN = lambda v: int(v * 914400)
L, W = MARGIN_L, CONTENT_W

SENTINEL_EN = "The rule, and how it was checked"
SENTINEL_AR = "القاعدة، وكيف تحقّقنا منها"

TEXT_EN = [
    ("We changed nothing in the maths after the audit. The two proposed fixes are shown in Part 3, and the decision is yours.",
     "Part 3 carries the APA rules, the measurement behind each constant, and the two weak points we are not hiding."),
    ("We audited the APA: the rule is right, but two questions can count twice",
     "The APA rules, the cap that keeps them honest, and where the mechanism is still weak"),
    ("A consequence worth stating: only FOUR scenarios can now move the profile, so the stability ceiling was re-measured from a fresh null model - it went from 65 down to 55.",
     "A consequence worth stating: only FOUR scenarios can move the profile, so the stability ceiling is measured from a null model over those four. It stands at 56."),
    ("It previously showed the participant's frozen score beside the company's, while the card below showed their LIVE score —\ntwo different numbers for the same person on one screen. \"The one you rated lowest\" could also become false.",
     "The participant's own priorities sit on screen immediately below it, so the comparison is theirs to draw — which is what the study measures.\nA number here would also have to choose between the frozen and the live profile, and either choice shows two figures for one person."),
    ("The EMPLOYER PRINCIPLE CARD rebuilt — slate letterhead, large quote, company's value in bold amber italic.\nIt now shows NO numbers.",
     "The EMPLOYER PRINCIPLE CARD — slate letterhead, large quote, company's value in bold amber italic.\nIt shows NO numbers."),
    ("The scenario page no longer opens part-way down.", "The scenario page always opens at its header."),
    ("The page behind the intro kept its own scrollbar, so reading the intro scrolled it. It is now frozen, and reset to top\nBEFORE the morph measures — a scrolled page would send the cards to the wrong place.",
     "The page behind the intro is frozen while it is up, and reset to its top BEFORE the morph measures anything —\nthe transition reads where each card sits in the viewport, so a scrolled page would send them to the wrong place."),
    ("Confidence moved to Question 2. \"How many are harmed\" renamed \"Reducing harm\" everywhere.",
     "Confidence sits with Question 2. The four values are named the same way everywhere."),
    ("The rating now sits under the value it scales. The arithmetic was NOT changed.\nThe old value name described the problem; the new one describes the choice.",
     "The rating sits directly under the value it scales.\n\"Reducing harm\" names what the participant is choosing rather than the problem they choose about."),
    ("What else changed since the last deck", "The Block 5 interface"),
    ("Interface and instrumentation. None of it touches a measure.", "None of it touches a measure."),
    ("Change", "Feature"),
]

TEXT_AR = [
    ("لم نغيّر شيئاً في الحسابات بعد التدقيق. الإصلاحان المقترحان معروضان في الجزء 3، والقرار قرارك.",
     "الجزء 3 يحمل قواعد APA، والقياس خلف كل ثابت، ونقطتَي الضعف اللتين لا نخفيهما."),
    ("دقّقنا آلية APA: القاعدة صحيحة، لكن سؤالين قد يُحتسبان مرتين",
     "قواعد APA، والسقف الذي يبقيها صادقة، ومواضع الضعف المتبقّية"),
    ("نتيجة تستحق الذكر: أربعة سيناريوهات فقط صارت قادرة على تحريك الصورة، لذلك أُعيد قياس سقف الثبات من نموذج عشوائي جديد — فانتقل من 65 إلى 55.",
     "نتيجة تستحق الذكر: أربعة سيناريوهات فقط تحرّك الصورة، لذا يُقاس سقف الثبات من نموذج عشوائي على تلك الأربعة. وهو 56."),
    ("ما الذي تغيّر أيضاً منذ العرض السابق", "واجهة البلوك الخامس"),
    ("واجهة وأدوات قياس. ولا شيء منها يمسّ مقياساً.", "لا شيء منها يمسّ مقياساً."),
    ("التغيير", "الميزة"),
]


def rebuild_opener(prs, ar):
    """The Part 3 opener, restated for the rule as it stands."""
    key = "The APA, audited" if not ar else "آلية APA — نتائج التدقيق"
    for s in prs.slides:
        heads = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if not heads or key not in heads[0]:
            continue
        page = None
        for sh in s.shapes:
            if (sh.has_text_frame and sh.left is not None
                    and abs(sh.top - D.PAGENO[1]) < 20000 and sh.text_frame.text.strip().isdigit()):
                page = sh.text_frame.text.strip()
        for sh in list(s.shapes):
            sh._element.getparent().remove(sh._element)
        if ar:
            head(s, SENTINEL_AR, "قاعدة واحدة، و19 اختباراً آلياً، وستة مشاركين يجيبون إجابات مختلفة.", page or 0)
            box(s, L, IN(1.85), W, IN(1.10),
                "القاعدة في جملة واحدة\n"
                "‏+30 للقيمة التي تسمّيها  ·  −20 للقيمة الأعلى حالياً  ·  الكل مضروب في 0.6 إلى 1.0 حسب تأكدك\n"
                "ولا تتحرّك أي قيمة أكثر من 30 × ذلك المعامل في التوضيح الواحد.",
                fill=ORANGE, size=15)
            box(s, L + int(W * 0.515), IN(3.15), int(W * 0.485), IN(1.15),
                "‏19 اختباراً آلياً تتحقق من هذه الجملة،\nوجميعها ناجحة.\n(npm run verify:apa)", fill=GREEN, size=14)
            box(s, L, IN(3.15), int(W * 0.485), IN(1.15),
                "وستة مشاركين بملفّ بداية متطابق\nوإجابات مختلفة، عبر الشيفرة الحقيقية.\n(npm run apa:personas)",
                fill=BLUE, size=14)
            box(s, L, IN(4.55), W, IN(1.30),
                "ما يلي: ما تسأله APA، والقواعد بالضبط، وماذا تفعل درجة الثقة، والمشاركون الستة،\n"
                "ولماذا يوجد سقف، وما يبقى ضعيفاً بعده.",
                fill=PANEL, size=14)
        else:
            head(s, SENTINEL_EN, "One rule, 19 automated checks, and six participants answering differently.", page or 0)
            box(s, L, IN(1.85), W, IN(1.10),
                "THE RULE, IN ONE SENTENCE\n"
                "+30 to the value you name  ·  −20 to the value currently on top  ·  all scaled 0.6–1.0 by how sure you say you are\n"
                "and no value moves more than 30 × that scale in one clarification.",
                fill=ORANGE, size=15)
            box(s, L, IN(3.15), int(W * 0.485), IN(1.15),
                "19 automated checks assert that sentence,\nand all of them pass.\n(npm run verify:apa)", fill=GREEN, size=14)
            box(s, L + int(W * 0.515), IN(3.15), int(W * 0.485), IN(1.15),
                "Six participants, identical starting profile,\ndifferent answers, through the real code.\n(npm run apa:personas)",
                fill=BLUE, size=14)
            box(s, L, IN(4.55), W, IN(1.30),
                "What follows: what the APA asks · the exact rules · what confidence is worth · the six participants\n"
                "· why there is a cap · and what stays weak in spite of it.",
                fill=PANEL, size=14)
        return True
    return False


def retext(prs, pairs):
    n = 0
    def walk(tf):
        nonlocal n
        whole = tf.text
        for a, b in pairs:
            if a not in whole:
                continue
            if "\n" in a:
                lines = b.split("\n")
                keep = tf.paragraphs[0]
                for extra in list(tf.paragraphs)[1:]:
                    extra._p.getparent().remove(extra._p)
                if keep.runs:
                    keep.runs[0].text = lines[0]
                    for r in list(keep.runs)[1:]:
                        r._r.getparent().remove(r._r)
                    for extra in lines[1:]:
                        np = tf.add_paragraph(); np.text = extra
                        np.alignment = keep.alignment; np.line_spacing = keep.line_spacing
                        kr = keep.runs[0]
                        for nr in np.runs:
                            nr.font.name = kr.font.name; nr.font.size = kr.font.size
                            nr.font.bold = kr.font.bold; nr.font.italic = kr.font.italic
                            try: nr.font.color.rgb = kr.font.color.rgb
                            except Exception: pass
                        if D._RTL and np._pPr is not None:
                            np._pPr.set("rtl", "1")
                    n += 1
            else:
                for para in tf.paragraphs:
                    for r in para.runs:
                        if a in r.text:
                            r.text = r.text.replace(a, b); n += 1
            whole = tf.text
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                walk(sh.text_frame)
            if sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        walk(c.text_frame)
    return n


def main(ar):
    name = "Block5_Design_Update_ARABIC" if ar else "Block5_Design_Update_for_Advisor"
    path = os.path.join(OUT, name + ".pptx")
    D.set_locale(font="Arial" if ar else "Calibri", rtl=ar)
    prs = Presentation(path)
    sentinel = SENTINEL_AR if ar else SENTINEL_EN
    if any(sentinel in sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame):
        print("REFUSING: %s already restated." % os.path.basename(path))
        return 1
    opener = rebuild_opener(prs, ar)
    swaps = retext(prs, TEXT_AR if ar else TEXT_EN)
    prs.save(path)
    print("OK  %s   opener rebuilt: %s   ·   text swaps: %d" % (os.path.basename(path), opener, swaps))
    return 0


if __name__ == "__main__":
    sys.exit(main("--ar" in sys.argv))
